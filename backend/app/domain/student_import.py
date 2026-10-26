"""Parse, validate, and commit bulk student imports.

Nothing is written to the student table until a previewed job is confirmed.
Confirm re-checks duplicates so rows created between preview and confirm are
skipped instead of failing the batch; any remaining race on the unique
constraint rolls the whole batch back.
"""
import csv
import io
import re
from dataclasses import dataclass
from datetime import UTC, date, datetime
from pathlib import Path
from uuid import UUID

from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.ext.asyncio import AsyncSession

from app.models import AuditEvent, Student, StudentImportJob, utcnow
from app.domain.subscriptions import require_student_capacity

REQUIRED_COLUMNS = ("admission_number", "first_name", "last_name")
OPTIONAL_COLUMNS = ("date_of_birth",)
TEMPLATE_COLUMNS = (*REQUIRED_COLUMNS, *OPTIONAL_COLUMNS)
MAX_IMPORT_ROWS = 2000
MAX_UPLOAD_BYTES = 2 * 1024 * 1024
SUPPORTED_EXTENSIONS = {".csv", ".xlsx"}

_TEMPLATE_SAMPLE = (
    ("AB-2026-001", "Ama", "Mensah", "2014-05-12"),
    ("AB-2026-002", "Kwame", "Asante", "2013-11-02"),
)


class ImportFileError(Exception):
    """The uploaded file itself cannot be used (type, headers, size, or empty)."""


class ImportConflict(Exception):
    """Confirmation cannot proceed (already confirmed, nothing to do, or a race)."""


@dataclass(frozen=True)
class ParsedRow:
    row_number: int
    admission_number: str
    first_name: str
    last_name: str
    date_of_birth: str | None


@dataclass
class RowOutcome:
    row_number: int
    admission_number: str
    first_name: str
    last_name: str
    date_of_birth: str | None
    status: str  # valid | error | duplicate | created
    message: str | None = None

    def as_dict(self) -> dict:
        return {
            "row_number": self.row_number,
            "admission_number": self.admission_number,
            "first_name": self.first_name,
            "last_name": self.last_name,
            "date_of_birth": self.date_of_birth,
            "status": self.status,
            "message": self.message,
        }


def build_template(extension: str) -> bytes:
    """Generate the downloadable import template for .csv or .xlsx."""
    extension = extension.lower().lstrip(".")
    if extension == "csv":
        buffer = io.StringIO()
        writer = csv.writer(buffer)
        writer.writerow(TEMPLATE_COLUMNS)
        writer.writerows(_TEMPLATE_SAMPLE)
        return buffer.getvalue().encode("utf-8")
    if extension == "xlsx":
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "students"
        sheet.append(list(TEMPLATE_COLUMNS))
        for sample in _TEMPLATE_SAMPLE:
            sheet.append(list(sample))
        buffer = io.BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()
    raise ImportFileError("Unsupported template format; use .csv or .xlsx")


def _normalize_header(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", "_", str(value or "").strip().lower()).strip("_")


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _split_rows(
    rows: list[list[object]],
) -> tuple[dict[str, int], dict[str, str], list[tuple[int, list[str]]]]:
    """Header row first: return normalized->index, normalized->label, data rows."""
    non_empty = [row for row in rows if any(_cell_text(cell) for cell in row)]
    if not non_empty:
        raise ImportFileError("The file is empty")
    header_cells = non_empty[0]
    index_map = {_normalize_header(name): index for index, name in enumerate(header_cells)}
    label_map = {_normalize_header(name): str(name).strip() for name in header_cells}
    data = [
        (row_number, [_cell_text(cell) for cell in cells])
        for row_number, cells in enumerate(non_empty[1:], start=2)
        if any(_cell_text(cell) for cell in cells)
    ]
    if not data:
        raise ImportFileError("The file has no data rows below the header")
    return index_map, label_map, data


def _read_csv(content: bytes) -> tuple[dict[str, int], dict[str, str], list[tuple[int, list[str]]]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as error:
        raise ImportFileError("The file must be a UTF-8 encoded CSV") from error
    reader = csv.reader(io.StringIO(text))
    try:
        rows = [list(row) for row in reader]
    except csv.Error as error:
        raise ImportFileError("The CSV could not be parsed") from error
    return _split_rows(rows)


def _read_xlsx(content: bytes) -> tuple[dict[str, int], dict[str, str], list[tuple[int, list[str]]]]:
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as error:
        raise ImportFileError("The file is not a readable Excel workbook") from error
    try:
        sheet = workbook.active
        if sheet is None:
            raise ImportFileError("The workbook has no worksheets")
        rows = [list(row) for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()
    return _split_rows(rows)


def parse_upload(filename: str, content: bytes) -> tuple[dict[str, str], list[ParsedRow]]:
    """Parse an uploaded file into mapped rows; raise ImportFileError when unusable."""
    if len(content) > MAX_UPLOAD_BYTES:
        raise ImportFileError(
            f"File is larger than the {MAX_UPLOAD_BYTES // (1024 * 1024)} MB import limit"
        )
    extension = Path(filename or "").suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise ImportFileError("Unsupported file type; upload a .csv or .xlsx file")
    index_map, label_map, data = _read_csv(content) if extension == ".csv" else _read_xlsx(content)
    if len(data) > MAX_IMPORT_ROWS:
        raise ImportFileError(
            f"File has {len(data)} rows; the limit is {MAX_IMPORT_ROWS} per import"
        )
    missing = [column for column in REQUIRED_COLUMNS if column not in index_map]
    if missing:
        found = ", ".join(sorted(label_map.values())) or "none"
        raise ImportFileError(
            f"Missing required column(s): {', '.join(missing)}. Found columns: {found}"
        )
    column_map = {column: label_map[column] for column in TEMPLATE_COLUMNS if column in label_map}
    parsed: list[ParsedRow] = []
    for row_number, cells in data:

        def cell(column: str, cells: list[str] = cells) -> str:
            index = index_map.get(column)
            if index is None or index >= len(cells):
                return ""
            return str(cells[index]).strip()

        parsed.append(
            ParsedRow(
                row_number=row_number,
                admission_number=cell("admission_number"),
                first_name=cell("first_name"),
                last_name=cell("last_name"),
                date_of_birth=cell("date_of_birth") or None,
            )
        )
    return column_map, parsed


def _validate_date(value: str | None) -> tuple[str | None, str | None]:
    if not value:
        return None, None
    text = value.strip()
    for pattern in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            parsed = datetime.strptime(text, pattern).date()
        except ValueError:
            continue
        if parsed > datetime.now(UTC).date():
            return None, "Date of birth cannot be in the future"
        return parsed.isoformat(), None
    return None, f"Date of birth must be YYYY-MM-DD or DD/MM/YYYY (received \"{value}\")"


def validate_rows(parsed: list[ParsedRow]) -> list[RowOutcome]:
    """Field-level validation plus in-file duplicate detection (no DB access)."""
    outcomes: list[RowOutcome] = []
    seen: dict[str, int] = {}
    for row in parsed:
        errors: list[str] = []
        admission = row.admission_number.strip().upper()
        if not admission:
            errors.append("Admission number is required")
        elif any(char.isspace() for char in admission):
            errors.append("Admission number cannot contain spaces")
        elif len(admission) > 40:
            errors.append("Admission number must be 40 characters or fewer")
        first_name = " ".join(row.first_name.split())
        last_name = " ".join(row.last_name.split())
        if not first_name:
            errors.append("First name is required")
        elif len(first_name) > 80:
            errors.append("First name must be 80 characters or fewer")
        if not last_name:
            errors.append("Last name is required")
        elif len(last_name) > 80:
            errors.append("Last name must be 80 characters or fewer")
        dob, dob_error = _validate_date(row.date_of_birth)
        if dob_error:
            errors.append(dob_error)
        duplicate_of = seen.get(admission) if admission else None
        if admission and duplicate_of is None:
            seen[admission] = row.row_number
        outcome = RowOutcome(
            row_number=row.row_number,
            admission_number=admission,
            first_name=first_name,
            last_name=last_name,
            date_of_birth=dob,
            status="valid",
        )
        if errors:
            outcome.status = "error"
            outcome.message = "; ".join(errors)
        elif duplicate_of is not None:
            outcome.status = "duplicate"
            outcome.message = f"Duplicate of row {duplicate_of} in this file"
        outcomes.append(outcome)
    return outcomes


async def flag_existing_duplicates(
    session: AsyncSession, school_id: UUID, outcomes: list[RowOutcome]
) -> None:
    """Mark valid outcomes whose admission number already exists at this school."""
    candidates = {item.admission_number for item in outcomes if item.status == "valid"}
    if not candidates:
        return
    existing = await session.execute(
        select(Student.admission_number).where(
            Student.school_id == school_id,
            Student.admission_number.in_(candidates),
        )
    )
    existing_ids = set(existing.scalars())
    for item in outcomes:
        if item.status == "valid" and item.admission_number in existing_ids:
            item.status = "duplicate"
            item.message = "Admission number already exists at this school"


def summarise(outcomes: list[RowOutcome]) -> dict[str, int]:
    """Partition rows into valid / duplicate / error counts (they sum to total)."""
    valid = sum(1 for item in outcomes if item.status == "valid")
    duplicates = sum(1 for item in outcomes if item.status == "duplicate")
    errors = sum(1 for item in outcomes if item.status == "error")
    return {
        "total_rows": len(outcomes),
        "valid_rows": valid,
        "duplicate_rows": duplicates,
        "error_rows": errors,
    }


async def commit_import(
    session: AsyncSession, school_id: UUID, actor_id: UUID, job: StudentImportJob
) -> dict[str, int]:
    """Create students for the previewed valid rows in a single transaction.

    Returns the final counts. Row statuses become created/duplicate/error, so
    created + duplicate + error always equals total_rows.
    """
    if job.status == "confirmed":
        raise ImportConflict("This import has already been confirmed")
    outcomes = [RowOutcome(**item) for item in job.preview_rows]
    preview_valid = sum(1 for item in outcomes if item.status == "valid")
    if preview_valid == 0:
        raise ImportConflict("This import has no valid rows to create")
    await flag_existing_duplicates(session, school_id, outcomes)
    importable = sum(1 for item in outcomes if item.status == "valid")
    await require_student_capacity(session, school_id, importable)
    created = 0
    for item in outcomes:
        if item.status != "valid":
            continue
        session.add(
            Student(
                school_id=school_id,
                admission_number=item.admission_number,
                first_name=item.first_name,
                last_name=item.last_name,
                date_of_birth=(
                    date.fromisoformat(item.date_of_birth) if item.date_of_birth else None
                ),
            )
        )
        item.status = "created"
        item.message = None
        created += 1
    try:
        await session.flush()
    except IntegrityError as error:
        # A row was inserted concurrently after our duplicate check; nothing is
        # committed, so the job stays previewed and can be retried safely.
        await session.rollback()
        raise ImportConflict(
            "A record was added at the same moment; no students were created. "
            "Preview the file again and retry"
        ) from error
    job.preview_rows = [item.as_dict() for item in outcomes]
    job.status = "confirmed"
    job.confirmed_at = utcnow()
    job.created_rows = created
    job.skipped_rows = preview_valid - created
    job.valid_rows = created
    job.duplicate_rows = sum(1 for item in outcomes if item.status == "duplicate")
    job.error_rows = sum(1 for item in outcomes if item.status == "error")
    session.add(
        AuditEvent(
            school_id=school_id,
            actor_user_id=actor_id,
            action="student.import_confirmed",
            resource_type="student_import",
            resource_id=job.id,
        )
    )
    await session.commit()
    return {
        "total_rows": job.total_rows,
        "valid_rows": job.valid_rows,
        "duplicate_rows": job.duplicate_rows,
        "error_rows": job.error_rows,
        "created_rows": created,
        "skipped_rows": job.skipped_rows,
    }


def error_csv(outcomes: list[dict]) -> bytes:
    """Render failed and duplicate rows as a downloadable CSV."""
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(
        ["row_number", "admission_number", "first_name", "last_name", "date_of_birth", "problem"]
    )
    for item in outcomes:
        if item.get("status") in ("error", "duplicate"):
            writer.writerow(
                [
                    item.get("row_number", ""),
                    item.get("admission_number", ""),
                    item.get("first_name", ""),
                    item.get("last_name", ""),
                    item.get("date_of_birth") or "",
                    item.get("message") or "",
                ]
            )
    return buffer.getvalue().encode("utf-8")
